"""Build deck/current-system.pptx — one slide per component question.

Regenerate with:  python deck/build_deck.py
Everything here is drawn from the code as it stands, not from the plan.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent / "current-system.pptx"

# ── the look ────────────────────────────────────────────────────────────────
SANS = "Helvetica Neue"
MONO = "Menlo"

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x76, 0x76, 0x76)
RULE = RGBColor(0xD9, 0xD7, 0xD1)
PANEL = RGBColor(0xF5, 0xF4, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

BLUE = RGBColor(0x2B, 0x5F, 0xD9)       # code / the system
BLUE_BG = RGBColor(0xE9, 0xEF, 0xFC)
AMBER = RGBColor(0xA9, 0x6B, 0x06)      # the model
AMBER_BG = RGBColor(0xFB, 0xF0, 0xDA)
GREEN = RGBColor(0x1B, 0x74, 0x50)      # a person
GREEN_BG = RGBColor(0xE3, 0xF2, 0xEB)
RED = RGBColor(0xA8, 0x25, 0x1E)        # refusal / stop
RED_BG = RGBColor(0xFA, 0xE9, 0xE6)

W, H = 13.333, 7.5


# ── primitives ──────────────────────────────────────────────────────────────
def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


def _plain(shape):
    shape.shadow.inherit = False
    return shape


def para(tf, text, *, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
         font=SANS, space_before=0, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(0)
    for chunk, strong in _segments(text):
        r = p.add_run()
        r.text = chunk
        r.font.size = Pt(size)
        r.font.bold = bold or strong
        r.font.color.rgb = color
        r.font.name = font
    return p


def _segments(text):
    """*emphasis* inside a string becomes a bold run."""
    out, buf, strong = [], "", False
    for part in text.split("*"):
        out.append((part, strong))
        strong = not strong
    return [(t, s) for t, s in out if t != ""] or [("", False)]


def text(slide, x, y, w, h, lines, *, size=12, color=INK, align=PP_ALIGN.LEFT,
         font=SANS, bold=False, spacing=4):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        para(tf, line, size=size, color=color, align=align, font=font, bold=bold,
             space_before=0 if i == 0 else spacing, first=(i == 0))
    return tb


def box(slide, x, y, w, h, lines, *, fill=WHITE, line=RULE, size=11.5, color=INK,
        bold=False, align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        font=SANS, line_w=1.0, spacing=3, radius=0.10):
    sh = _plain(slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.09)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        para(tf, ln, size=size, color=color, align=align, bold=bold, font=font,
             space_before=0 if i == 0 else spacing, first=(i == 0))
    return sh


def chip(slide, x, y, w, h, label, *, fg=BLUE, bg=BLUE_BG, size=10):
    return box(slide, x, y, w, h, label, fill=bg, line=None, size=size, color=fg,
               bold=True, radius=0.5)


def line(slide, x1, y1, x2, y2, *, color=RULE, w=1.0, arrow=True, dash=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(w)
    ln = cn.line._get_or_add_ln()
    if dash:
        d = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
        ln.append(d)
    if arrow:
        tail = ln.makeelement(qn("a:tailEnd"),
                              {"type": "triangle", "w": "med", "len": "med"})
        ln.append(tail)
    return cn


def rule(slide, x, y, w, *, color=RULE, weight=1.0):
    line(slide, x, y, x + w, y, color=color, w=weight, arrow=False)


PAGE = {"n": 0}


def slide(prs, kicker, question, *, accent=BLUE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    PAGE["n"] += 1
    text(s, 0.62, 0.40, 10.0, 0.25, kicker.upper(), size=10, color=accent, bold=True)
    text(s, 0.62, 0.68, 11.8, 0.55, question, size=25, color=INK, bold=True)
    rule(s, 0.62, 1.38, 12.1)
    text(s, 12.15, 6.95, 0.6, 0.25, str(PAGE["n"]), size=9.5, color=MUTED,
         align=PP_ALIGN.RIGHT)
    return s


def footnote(s, note):
    text(s, 0.62, 6.93, 10.5, 0.3, note, size=9.5, color=MUTED)


# ═══════════════════════════════════════════════════════════════════════════
prs = deck()

# ── 1. title ────────────────────────────────────────────────────────────────
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
box(s, 0, 0, 0.30, H, "", fill=BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.05, 1.55, 10, 0.3, "COMPUTER-USE AUTOMATION  ·  BUILD REVIEW", size=11,
     color=BLUE, bold=True)
text(s, 1.02, 2.00, 11.2, 1.0, "What the system actually looks like", size=40, color=INK,
     bold=True)
text(s, 1.05, 2.95, 10.6, 1.0,
     "An LLM discovers a task on a legacy bank app *once*. The run is compiled into a "
     "reviewable *Artifact*. Thereafter a deterministic *Replay Engine* executes it — "
     "with no model in the decision loop.", size=15, color=MUTED, spacing=6)
rule(s, 1.05, 4.25, 11.0)

facts = [("13", "modules in cua/", BLUE, BLUE_BG),
         ("98", "tests, no mocking", GREEN, GREEN_BG),
         ("4", "policy layers, 4 owners", AMBER, AMBER_BG),
         ("2", "tenants, 1 artifact", BLUE, BLUE_BG)]
for i, (n, label, fg, bg) in enumerate(facts):
    x = 1.05 + i * 2.80
    box(s, x, 4.60, 2.50, 1.05, "", fill=bg, line=None)
    text(s, x + 0.14, 4.70, 1.0, 0.4, n, size=26, color=fg, bold=True)
    text(s, x + 0.14, 5.16, 2.2, 0.4, label, size=10.5, color=fg)

text(s, 1.05, 6.20, 11, 0.3,
     "Each slide answers one question about one component. Comparison with the original "
     "sketch is on the last two slides.", size=11.5, color=MUTED)

# ── 2. the lifecycle ────────────────────────────────────────────────────────
s = slide(prs, "Lifecycle", "How does a task become automation?")

stages = [
    ("1  Goal + Contract", "A Reviewer fixes typed inputs,\noutputs and outcome codes\n*before* discovery runs.",
     GREEN, GREEN_BG, "Reviewer"),
    ("2  Discovery Run", "The LLM proposes one action\nper turn against a *non-prod*\ncopy. Code performs them.",
     AMBER, AMBER_BG, "Discovery LLM"),
    ("3  Recorder", "Turns the trace into a *draft*\nArtifact. Guesses are attached\nas suggestions, not decisions.",
     BLUE, BLUE_BG, "code"),
    ("4  Review + Approve", "Decisions file applied; lint +\na verify-replay on *unseen*\ninputs must pass.",
     GREEN, GREEN_BG, "2 Reviewers"),
    ("5  Replay", "The approved Artifact runs\nwith *no model at all*. The only\nway it runs in production.",
     BLUE, BLUE_BG, "Replay Engine"),
]
bw, gap = 2.18, 0.30
for i, (title, body, fg, bg, owner) in enumerate(stages):
    x = 0.62 + i * (bw + gap)
    box(s, x, 1.95, bw, 0.46, title, fill=fg, line=None, size=11.5, color=WHITE, bold=True)
    box(s, x, 2.41, bw, 1.42, body.split("\n"), fill=bg, line=None, size=10.5, color=INK,
        align=PP_ALIGN.LEFT, spacing=2)
    text(s, x, 3.92, bw, 0.25, owner, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 4:
        line(s, x + bw + 0.04, 2.18, x + bw + gap - 0.04, 2.18, color=MUTED, w=1.25)

box(s, 0.62, 4.55, 5.85, 1.25,
    ["*The artifact is data, not generated code.*",
     "It names actions, predicates and target-finders from closed vocabularies the engine "
     "already holds. An artifact cannot express anything else — so no artifact can bypass "
     "a guarantee (ADR 0001)."],
    fill=PANEL, line=None, size=11, align=PP_ALIGN.LEFT, spacing=5)

box(s, 6.87, 4.55, 5.85, 1.25,
    ["*Discovery happens once; replay happens forever.*",
     "One real discovery run: 13 turns, goal reached. Replay of the same capability then "
     "passes all 8 scenario tests deterministically — same flow, no tokens, no variance."],
    fill=PANEL, line=None, size=11, align=PP_ALIGN.LEFT, spacing=5)

footnote(s, "cua/discovery.py · cua/recorder.py · cua/review.py · cua/engine.py")

# ── 3. discovery ────────────────────────────────────────────────────────────
s = slide(prs, "Discovery Run", "What happens on each turn with the model?", accent=AMBER)

lanes = [("Browser / Surface", 1.55, BLUE), ("Engine + Policy", 5.55, BLUE), ("Discovery LLM", 10.30, AMBER)]
for label, x, fg in lanes:
    box(s, x - 1.25, 1.72, 2.50, 0.40, label, fill=fg, line=None, size=11, color=WHITE, bold=True)
    line(s, x, 2.16, x, 5.45, color=RULE, w=1.25, arrow=False, dash=True)

msgs = [
    (5.55, 10.30, 2.55, "observation: accessibility list + screenshot, values default-denied", AMBER),
    (10.30, 5.55, 3.10, "one tool call:  type → [4], secret=login_password", AMBER),
    (5.55, 1.55, 4.20, "perform — secret substituted at the keystroke", BLUE),
    (1.55, 5.55, 4.75, "the new screen", BLUE),
]
for x1, x2, y, label, col in msgs:
    line(s, x1, y, x2, y, color=col, w=1.5)
    text(s, min(x1, x2) + 0.12, y - 0.30, abs(x2 - x1) - 0.24, 0.26, label, size=10.5,
         color=INK, align=PP_ALIGN.CENTER)

box(s, 6.10, 3.30, 3.35, 0.50, "policy gate: origin · route · action type",
    fill=BLUE_BG, line=None, size=10.5, color=BLUE, bold=True)
text(s, 6.10, 3.85, 3.35, 0.25, "refused → the model is told why, and the turn is spent",
     size=9.5, color=RED, align=PP_ALIGN.CENTER)

text(s, 0.62, 5.05, 2.0, 0.3, "loop, up to 24 turns", size=10, color=MUTED)
line(s, 1.55, 5.45, 10.30, 5.45, color=RULE, w=1.0, arrow=False)

box(s, 0.62, 5.72, 6.05, 1.10,
    ["*Six endings, and the model only picks four.*",
     "goal_reached · report_outcome · ask_human · give_up   —   plus step-limit/timeout and "
     "stuck-detected, which code decides."],
    fill=PANEL, line=None, size=10.5, align=PP_ALIGN.LEFT, spacing=4)
box(s, 6.87, 5.72, 5.85, 1.10,
    ["*What the model never gets.*",
     "A secret value · a production environment · a field value whose Target is not a declared "
     "Readable Region · the ability to act outside the allowlist."],
    fill=PANEL, line=None, size=10.5, align=PP_ALIGN.LEFT, spacing=4)

footnote(s, "cua/discovery.py — bounds: 24 turns, 300 s, 60 actions, 3 consecutive failures")

# ── 4. the artifact ─────────────────────────────────────────────────────────
s = slide(prs, "Artifact", "What exactly is recorded?")

nodes = ["s1_login", "s2_search", "s4_member", "s5_form", "s6_review", "s7_done"]
nw = 1.58
for i, n in enumerate(nodes):
    x = 0.62 + i * (nw + 0.42)
    last = i == len(nodes) - 1
    box(s, x, 1.85, nw, 0.52, n, fill=GREEN_BG if last else WHITE,
        line=GREEN if last else RULE, size=10.5, color=GREEN if last else INK,
        bold=last, font=MONO)
    if i < 5:
        consequential = i == 4
        line(s, x + nw + 0.03, 2.11, x + nw + 0.39, 2.11,
             color=RED if consequential else MUTED, w=1.75 if consequential else 1.25)

text(s, 0.62, 2.45, 3.4, 0.25, "each State carries a *Checkpoint*", size=10, color=MUTED)
text(s, 8.90, 2.45, 3.9, 0.5,
     "this edge is *consequential* — it holds a\nVerification Check", size=10, color=RED,
     spacing=2)

parts = [
    ("Contract", "3 typed inputs · 2 outputs ·\n2 outcome codes. Fixed first,\nnever inferred from a run.", GREEN),
    ("States + Transitions", "6 states, 11 transitions.\nActions from a closed set:\nclick · type · select · read.", BLUE),
    ("Targets", "12 named controls, each with\nan ordered ladder of ways to\nfind it. Never a raw selector.", BLUE),
    ("Watchers", "3 here + 4 shared by every\nartifact on this app. Each one\ncarries its provenance.", AMBER),
    ("Needs", "The pages, actions and secrets\nthe run actually used. Checked\nagainst policy before replay.", RED),
]
pw = 2.34
for i, (title, body, fg) in enumerate(parts):
    x = 0.62 + i * (pw + 0.19)
    box(s, x, 3.30, pw, 0.40, title, fill=fg, line=None, size=11, color=WHITE, bold=True)
    box(s, x, 3.70, pw, 1.35, body.split("\n"), fill=PANEL, line=None, size=10,
        align=PP_ALIGN.LEFT, spacing=2)

box(s, 0.62, 5.45, 12.10, 1.30,
    ["*A checkpoint never contains a discovery run's literal values, and a lint fails the "
     "build if one survives.*",
     "The run that recorded this used member 54321 and nickname “Holiday fund”. The artifact "
     "holds {{member_number}} and {{nickname}}. Two separate defences: the *schema* is a "
     "whitelist, so an action like `download` cannot be parsed at all; the *lint* catches "
     "artifacts that are well-formed and still wrong — a stray literal, an unreachable outcome, "
     "needs outside the declared role, a consequential step with no verification check."],
    fill=WHITE, line=RULE, size=10.5, align=PP_ALIGN.LEFT, spacing=5)

footnote(s, "artifacts/open_sub_account.1.0.0.yaml · cua/artifact.py · cua/lint.py")

# ── 5. targeting ────────────────────────────────────────────────────────────
s = slide(prs, "Targeting", "How does it find a control on a page with no ids?")

rungs = [
    ("1", "role_name", "the accessibility tree's\ncomputed name",
     "moving, restyling", "a rename", BLUE),
    ("2", "label_anchor", "visible words plus layout:\n“the textbox right of *Member number*”",
     "a control being renamed", "the caption moving", GREEN),
    ("3", "picture", "a crop saved at record time",
     "renames and re-layout", "re-skinning", AMBER),
]
for i, (n, name, reads, survives, breaks, fg) in enumerate(rungs):
    y = 1.85 + i * 1.02
    box(s, 0.62, y, 0.52, 0.80, n, fill=fg, line=None, size=15, color=WHITE, bold=True)
    box(s, 1.20, y, 2.35, 0.80, name, fill=WHITE, line=fg, size=12, color=fg, bold=True, font=MONO)
    text(s, 3.70, y + 0.10, 2.95, 0.6, reads, size=10.5, color=INK, spacing=2)
    text(s, 6.80, y + 0.10, 2.35, 0.6, ["survives", survives], size=10.5, color=GREEN, spacing=2)
    text(s, 9.35, y + 0.10, 3.30, 0.6, ["breaks on", breaks], size=10.5, color=RED, spacing=2)
    if i < 2:
        line(s, 0.88, y + 0.83, 0.88, y + 0.99, color=MUTED, w=1.25)

text(s, 3.70, 1.60, 9.0, 0.25,
     "Tried in order. Replay records which rung matched — and stops rather than guessing when none do.",
     size=10, color=MUTED)

box(s, 0.62, 5.05, 5.90, 1.70,
    ["*The demo this exists for.*",
     "The search control is an `<img>` inside a `<button>` with no alt and no label, so the "
     "browser computes *no accessible name*: `get_by_role(\"button\", name=\"Search\")` returns "
     "0 matches. Its recorded ladder has no rung 1 at all, because there was nothing to record. "
     "label_anchor “Member number” finds it. *A text-only agent cannot see this control.*"],
    fill=RED_BG, line=None, size=10.5, align=PP_ALIGN.LEFT, spacing=5)

box(s, 6.82, 5.05, 5.90, 1.70,
    ["*A Target stores a relationship, never a measurement.*",
     "Distances are recomputed from the live page every run, so a re-layout does not "
     "invalidate the recording.",
     "Any match below rung 1 is a *Fallback Match*: allowed, always logged, and a rising rate "
     "for one tenant is the signal that its app has drifted."],
    fill=PANEL, line=None, size=10.5, align=PP_ALIGN.LEFT, spacing=5)

footnote(s, "cua/surface.py · docs/targeting.md · tools/demos/b1.py")

# ── 6. replay step ──────────────────────────────────────────────────────────
s = slide(prs, "Replay Engine", "What does replay do for a single step?")

steps = [
    ("1", "Verify before acting", "does the *from* state's checkpoint hold?", "surprise handler →", AMBER),
    ("2", "Resolve the Target", "walk the ladder; record the rung that matched", "Failed  target_not_found", RED),
    ("3", "Policy check, here and now", "this route, this action type", "Failed  policy_denied", RED),
    ("4", "Risk gate", "consequential actions are logged and gated", "Refused  at the front door", RED),
    ("5", "Act", "click · type · select · read — and nothing else", None, BLUE),
    ("6", "Observe", "does the *to* state's checkpoint hold?", "surprise handler →", AMBER),
]
for i, (n, title, body, fail, fg) in enumerate(steps):
    y = 1.72 + i * 0.86
    box(s, 0.62, y, 0.48, 0.66, n, fill=BLUE if fail is None or fg == BLUE else INK,
        line=None, size=13, color=WHITE, bold=True)
    box(s, 1.16, y, 5.55, 0.66, [f"*{title}*", body], fill=WHITE, line=RULE, size=10.5,
        align=PP_ALIGN.LEFT, spacing=3)
    if i < 5:
        line(s, 0.86, y + 0.69, 0.86, y + 0.84, color=MUTED, w=1.25)
    if fail:
        line(s, 6.78, y + 0.33, 7.45, y + 0.33, color=fg, w=1.25)
        box(s, 7.50, y + 0.08, 3.25, 0.50, fail, fill=RED_BG if fg == RED else AMBER_BG,
            line=None, size=10.5, color=fg, bold=True, align=PP_ALIGN.LEFT)

box(s, 11.00, 1.72, 1.72, 5.00,
    ["*No model.*", "", "Replay calls no LLM at all.",
     "", "Every decision is a predicate over the screen, evaluated by code.",
     "", "Same inputs, same path, every time."],
    fill=BLUE_BG, line=None, size=10.5, color=BLUE, align=PP_ALIGN.LEFT, spacing=1)

footnote(s, "cua/engine.py — 11 transitions, a loop guard, and a write-ahead record of every step")

# ── 7. surprises ────────────────────────────────────────────────────────────
s = slide(prs, "Error handling", "What happens when the screen is not what was expected?",
          accent=AMBER)

box(s, 0.62, 1.70, 3.00, 0.64, "checkpoint missed", fill=AMBER, line=None, size=12,
    color=WHITE, bold=True)
line(s, 3.68, 2.02, 4.30, 2.02, color=MUTED, w=1.25)
box(s, 4.35, 1.70, 3.30, 0.64, "does a *Watcher* recognise\nthis screen?".split("\n"),
    fill=WHITE, line=AMBER, size=11, align=PP_ALIGN.CENTER, spacing=1)

conds = [
    ("Business Outcome", "the app worked and gave a\nlegitimate non-happy answer",
     "stop — never retried", GREEN),
    ("Recoverable", "a dismissible or transient\ncondition, bounded budget",
     "the system fixes it itself", BLUE),
    ("Escalate", "only a person can fix it\nduring this run",
     "hand the live session over", AMBER),
    ("Hard Failure", "broken in a way retrying\nwill not fix",
     "stop with evidence", RED),
]
cw = 2.95
for i, (title, body, react, fg) in enumerate(conds):
    x = 0.62 + i * (cw + 0.19)
    box(s, x, 2.90, cw, 0.42, title, fill=fg, line=None, size=11.5, color=WHITE, bold=True)
    box(s, x, 3.32, cw, 1.10, body.split("\n") + ["", react], fill=PANEL, line=None,
        size=10, align=PP_ALIGN.LEFT, spacing=1)
    line(s, 6.00, 2.40, x + cw / 2, 2.86, color=RULE, w=1.0)

box(s, 8.00, 1.70, 4.72, 0.98,
    ["no watcher matches → *Unknown State*",
     "Escalated when attended, Failed when not. *Never guessed through.*"],
    fill=RED_BG, line=None, size=10.5, color=RED, align=PP_ALIGN.LEFT, spacing=4)

text(s, 0.62, 4.70, 6.0, 0.25, "The test is *who can act*: the system alone, a person during "
     "the run, or nobody in time.", size=10.5, color=MUTED)

rule(s, 0.62, 5.15, 12.10)
text(s, 0.62, 5.32, 4.0, 0.25, "The caller receives exactly one Run Result", size=11.5,
     color=INK, bold=True)
results = [("Succeeded", GREEN, GREEN_BG), ("Business Outcome", GREEN, GREEN_BG),
           ("Failed", RED, RED_BG), ("Aborted", MUTED, PANEL),
           ("Refused", RED, RED_BG), ("Outcome Unknown", AMBER, AMBER_BG)]
for i, (name, fg, bg) in enumerate(results):
    chip(s, 0.62 + i * 2.06, 5.72, 1.92, 0.44, name, fg=fg, bg=bg, size=10.5)

box(s, 0.62, 6.30, 12.10, 0.52,
    ["*Outcome Unknown* means a consequential action was performed but its effect could "
     "never be confirmed — the one result that says *do not retry, a person must look*. "
     "The verification check exists so the system looks instead of clicking again."],
    fill=WHITE, line=RULE, size=10.5, align=PP_ALIGN.LEFT)

footnote(s, "docs/error-taxonomy.md · cua/engine.py::_handle_surprise")

# ── 8. policy ───────────────────────────────────────────────────────────────
s = slide(prs, "Policy", "Who decides what the automation may do?", accent=RED)

layers = [
    ("Baseline", "the provider (us)", "config/baseline.yaml",
     "4 allowed actions, 9 named\nand denied. Schemes, denied\nroute keywords, hard rules.", BLUE),
    ("Role", "per vendor app,\nwritten before discovery", "config/roles/",
     "pages · action types · secrets ·\nmay it act consequentially ·\nwhich service account.", AMBER),
    ("Tenant grant", "the institution", "config/policies/",
     "which roles it grants at all,\nits origin, its environment tag,\nany further narrowing.", GREEN),
    ("Needs", "derived from the run", "in the artifact",
     "the pages, actions and secrets\nthe discovery run actually\nused. Replay only.", RED),
]
lw = 2.80
for i, (title, owner, where, body, fg) in enumerate(layers):
    x = 0.62 + i * (lw + 0.35)
    box(s, x, 1.80, lw, 0.44, title, fill=fg, line=None, size=12.5, color=WHITE, bold=True)
    text(s, x + 0.02, 2.30, lw, 0.45, owner, size=10, color=MUTED, spacing=1)
    box(s, x, 2.82, lw, 1.28, body.split("\n"), fill=PANEL, line=None, size=10,
        align=PP_ALIGN.LEFT, spacing=2)
    text(s, x + 0.02, 4.15, lw, 0.25, where, size=9.5, color=fg, font=MONO)
    if i < 3:
        text(s, x + lw + 0.02, 1.86, 0.33, 0.35, "∩", size=19, color=INK,
             align=PP_ALIGN.CENTER, bold=True)

text(s, 0.62, 1.52, 9, 0.25,
    "Four files, four different owners — and each layer can only *narrow* the one before it.",
    size=10.5, color=MUTED)

box(s, 0.62, 4.62, 5.90, 1.05,
    ["*Refused at the front door, before a browser opens.*",
     "bank_b does not grant `account_opener`, so the same artifact that works for bank_a is "
     "refused — and the refusal names the layer that refused it."],
    fill=RED_BG, line=None, size=10.5, align=PP_ALIGN.LEFT, spacing=4)

box(s, 6.82, 4.62, 5.90, 1.05,
    ["*Enforced on every request the browser makes.*",
     "A check before we act cannot see a request the *page* starts. `/leaky` renders a 1×1 "
     "image pointing off-origin with member data in the query string. Nobody clicks; the "
     "allowlist aborts it anyway."],
    fill=RED_BG, line=None, size=10.5, align=PP_ALIGN.LEFT, spacing=4)

box(s, 0.62, 5.82, 12.10, 0.95,
    ["*The least-privilege chain ends in a real login.*",
     "The Service Account follows from the Role, so the credential is a consequence of the "
     "capability's intent. A `balance_reader` signs in as `svc_read`, which *the bank application "
     "itself* refuses when it reaches the sub-account form — so the guarantee survives every "
     "layer above it failing. Consequential actions additionally need *two named reviewers*."],
    fill=WHITE, line=RULE, size=10.5, align=PP_ALIGN.LEFT, spacing=4)

footnote(s, "cua/policy.py · ADR 0005 — effective permissions = Baseline ∩ Role ∩ Tenant grant ∩ Needs")

# ── 9. redaction ────────────────────────────────────────────────────────────
s = slide(prs, "Sensitive data", "How is member data kept out of the model, the logs and the artifact?",
          accent=RED)

text(s, 0.62, 1.58, 5.6, 0.5,
     "A name is just words; a birthday is just digits. *No pattern can find them.* So the first "
     "test is not what a value looks like — it is where it came from.", size=11, color=INK,
     spacing=3)

flow = [("a value on screen", WHITE, RULE, INK),
        ("is it a password?", WHITE, AMBER, INK),
        ("is its Target a declared\n*Readable Region*?", WHITE, AMBER, INK),
        ("pattern net\nssn · card · email · phone · date · money", WHITE, BLUE, INK),
        ("recorded", GREEN_BG, GREEN, GREEN)]
outs = [None, ("yes → (protected)   never read at all", RED),
        ("no → (hidden)   catches names, birthdays, and\nevery field nobody has thought about yet", RED),
        None, None]
for i, (label, bg, ln, col) in enumerate(flow):
    y = 1.98 + i * 0.78
    box(s, 0.62, y, 3.55, 0.64, label.split("\n"), fill=bg, line=ln, size=10.5, color=col,
        bold=(i == 4), spacing=1)
    if i < 4:
        line(s, 2.40, y + 0.67, 2.40, y + 0.81, color=MUTED, w=1.25)
    if outs[i]:
        txt, c = outs[i]
        line(s, 4.23, y + 0.32, 4.80, y + 0.32, color=c, w=1.25)
        text(s, 4.88, y + 0.12, 3.0, 0.5, txt, size=9.5, color=c, spacing=1)

rows = [
    ("evidence and logs", "only identifiers *we* generated are written", "total", GREEN),
    ("artifact", "example values became placeholders; a lint enforces it", "total", GREEN),
    ("secrets", "substituted at the keystroke, below the model and the log", "total", GREEN),
    ("outputs to the caller", "returned in full — they are the answer — masked in the record", "total", GREEN),
    ("observation to a model", "default-deny by Readable Region, then patterns", "declared fields", GREEN),
    ("watcher extraction", "a declared capture group only, never free page text", "total", GREEN),
    ("screenshots", "cropped, downscaled, declared Sensitive Regions painted black", "partial", RED),
]
text(s, 8.10, 1.58, 4.6, 0.3, "Every channel out, and what it guarantees", size=11.5,
     color=INK, bold=True)
for i, (chan, gate, guarantee, fg) in enumerate(rows):
    y = 2.00 + i * 0.62
    bg = PANEL if i % 2 == 0 else WHITE
    box(s, 8.10, y, 4.62, 0.58, "", fill=bg, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 8.22, y + 0.05, 2.0, 0.25, chan, size=10, color=INK, bold=True)
    text(s, 8.22, y + 0.26, 3.4, 0.28, gate, size=9, color=MUTED)
    text(s, 11.70, y + 0.15, 0.95, 0.25, guarantee, size=9, color=fg, bold=True,
         align=PP_ALIGN.RIGHT)

box(s, 0.62, 5.95, 7.10, 0.80,
    ["*The two channels take opposite defaults, deliberately.* Values are an allowlist — hiding "
     "one costs nothing. Pixels are a deny-list — blacking out a control the model must click "
     "would blind it. The image channel is the weaker one, and is documented as the residual risk."],
    fill=RED_BG, line=None, size=10, align=PP_ALIGN.LEFT)

footnote(s, "cua/redact.py · docs/security-model.md · 17 PII tests, including a name and a birthday no regex could catch")

# ── 10. handoff ─────────────────────────────────────────────────────────────
s = slide(prs, "Human handoff", "What happens when only a person can unblock the run?",
          accent=GREEN)

lease = [("automation", BLUE), ("awaiting_operator", AMBER), ("operator_in_control", GREEN),
         ("resuming", AMBER), ("automation", BLUE)]
labels = ["pause", "take", "resume", "re-orient"]
lwid = 2.28
for i, (name, fg) in enumerate(lease):
    x = 0.62 + i * (lwid + 0.42)
    box(s, x, 2.05, lwid, 0.66, name, fill=fg, line=None, size=11.5, color=WHITE,
        bold=True, font=MONO)
    if i < 4:
        line(s, x + lwid + 0.03, 2.38, x + lwid + 0.39, 2.38, color=MUTED, w=1.25)
        text(s, x + lwid, 2.44, 0.45, 0.25, labels[i], size=9, color=MUTED,
             align=PP_ALIGN.CENTER)

text(s, 0.62, 1.70, 9.0, 0.25,
     "Control is a *lease*: one holder at a time, and the illegal moves raise rather than "
     "being prevented by convention.", size=10.5, color=MUTED)
text(s, 0.62, 2.82, 12.1, 0.25, "any state may also move to  *done*  — the Operator can abort, "
     "and an abort is a Run Result, not a failure", size=9.5, color=MUTED)

cards = [
    ("It is the same session", "The Operator gets the browser the automation was already "
     "using — mid-flow, on the same page — not a fresh one.", GREEN),
    ("They arrive with context", "An intervention record: the capability, the state, the "
     "reason, the URL and a screenshot.", BLUE),
    ("What they did is recorded;\nwhat they typed is not", "Decision and whether they "
     "navigated are written to the trail. A test greps it for the password.", RED),
    ("Resume does not assume\nthey finished", "The engine asks which checkpoint holds. "
     "An Operator who wandered off gets resume_checkpoint_missed, not a run "
     "carrying on in the wrong place.", AMBER),
]
cwid = 3.00
for i, (title, body, fg) in enumerate(cards):
    x = 0.62 + i * (cwid + 0.19)
    box(s, x, 3.40, cwid, 0.62, title.split("\n"), fill=fg, line=None, size=11,
        color=WHITE, bold=True, spacing=1)
    box(s, x, 4.02, cwid, 1.55, body, fill=PANEL, line=None, size=10.5,
        align=PP_ALIGN.LEFT)

box(s, 0.62, 5.85, 12.10, 0.85,
    ["*Escalations are bounded at two per state*, so escalate → resume → escalate cannot keep a "
     "person answering the same question forever — the third time is a Run Result, not another "
     "interruption. And the Operator is the *only* human who ever touches a live session: a "
     "Reviewer reads redacted evidence, afterwards (ADR 0002)."],
    fill=WHITE, line=RULE, size=10.5, align=PP_ALIGN.LEFT)

footnote(s, "cua/handoff.py · cua/engine.py::_hand_over · tests/test_handoff.py")

# ── 11. overlays ────────────────────────────────────────────────────────────
s = slide(prs, "Multi-tenant", "How does one recording serve banks that look different?")

box(s, 4.55, 1.75, 4.25, 1.00,
    ["Artifact  member.open_sub_account@1.0.0", "recorded at First Credit Union"],
    fill=BLUE, line=None, size=11.5, color=WHITE, bold=True, spacing=3)

line(s, 5.60, 2.78, 4.10, 3.45, color=MUTED, w=1.25)
line(s, 7.75, 2.78, 9.25, 3.45, color=MUTED, w=1.25)

box(s, 2.20, 3.50, 3.60, 0.92, ["First Credit Union", "runs unchanged"],
    fill=GREEN_BG, line=None, size=11, color=GREEN, bold=True, spacing=3)
box(s, 7.55, 3.50, 3.60, 0.92, ["Lakeside Savings", "same product, rebranded"],
    fill=AMBER_BG, line=None, size=11, color=AMBER, bold=True, spacing=3)

box(s, 7.55, 4.58, 3.60, 1.42,
    ["*+ Tenant Overlay  (20 lines)*",
     "“Member number” → “Find member by #”",
     "the search icon moved to the *left*, so its rung needs relation `nearest`",
     "“Open” → “Create”,  “Continue” → “Submit”"],
    fill=WHITE, line=AMBER, size=9.5, align=PP_ALIGN.LEFT, spacing=4)
line(s, 9.35, 4.44, 9.35, 4.54, color=MUTED, w=1.25)

box(s, 0.62, 4.58, 6.55, 1.42,
    ["*Without the overlay it fails honestly*, at the first checkpoint, with unknown_state — "
     "because the field is called something else there. It does not guess its way forward."],
    fill=PANEL, line=None, size=10.5, align=PP_ALIGN.LEFT)

box(s, 0.62, 6.15, 12.10, 0.62,
    ["*An overlay may change how things look, never how the capability behaves.* Origin, labels, "
     "anchors, positions, timeouts — yes. Transitions, contract or safety policy — refused by "
     "the lint before the browser opens. If a tenant needs different behaviour, that is a "
     "different artifact."],
    fill=RED_BG, line=None, size=10.5, align=PP_ALIGN.LEFT)

footnote(s, "overlays/lakeside.yaml · cua/overlay.py · tools/demos/b2.py")

# ── 12. delta ───────────────────────────────────────────────────────────────
s = slide(prs, "Then and now", "What changed from the original sketch — and why?")

text(s, 0.62, 1.58, 5.8, 0.25, "ORIGINAL SKETCH", size=10, color=MUTED, bold=True)
text(s, 6.95, 1.58, 5.8, 0.25, "WHAT IT BECAME", size=10, color=BLUE, bold=True)

deltas = [
    ("Record → LLM schema → replay contract → *Code*",
     "The recording stays *data*. Generated code could express anything; an artifact can only "
     "name things the engine already holds, so guarantees cannot be bypassed."),
    ("“All checked against guardrails / policies” — one gate",
     "*Four layers with four owners*, intersected: Baseline ∩ Role ∩ Tenant grant ∩ Needs. Plus "
     "interception of every request the page itself starts."),
    ("“Credential and PII need to be redacted”",
     "Redaction is decided by *origin*, not by pattern — default-deny by Readable Region, "
     "because no regex finds a name. Patterns are the second net only."),
    ("Error-handling table, “needs work”",
     "*Four Conditions* by who can act, *six Run Results*, and Watchers carrying provenance. "
     "Unknown State is never guessed through."),
    ("Human-in-the-loop, listed per exception",
     "One *control lease* with declared transitions, on the live session, bounded at two "
     "escalations per state."),
    ("state / action / transition, from the PreAct paper",
     "Kept — plus a *Checkpoint* on every state, a *Target ladder* per control, and a "
     "*Verification Check* on every consequential edge."),
]
for i, (before, after) in enumerate(deltas):
    y = 1.92 + i * 0.82
    if i % 2 == 0:
        box(s, 0.62, y - 0.06, 12.10, 0.78, "", fill=PANEL, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.75, y + 0.06, 5.55, 0.65, before, size=10.5, color=MUTED, spacing=2)
    line(s, 6.48, y + 0.32, 6.82, y + 0.32, color=BLUE, w=1.25)
    text(s, 6.95, y + 0.02, 5.75, 0.72, after, size=10.5, color=INK, spacing=2)

footnote(s, "The sequence of the discovery loop, and the state/action/transition shape, survived the build unchanged.")

# ── 13. status ──────────────────────────────────────────────────────────────
s = slide(prs, "Status", "What is proven, and what is not?", accent=GREEN)

proven = [
    ("98 tests, no mocking", "Scenario tests run against the live Flask app: the member "
     "number selects the scenario, so a test reads “replay 99999, expect MEMBER_NOT_FOUND”."),
    ("A real discovery run", "13 turns, goal reached, outputs correct. The model signed in "
     "with secrets it never saw, found the unlabelled icon by anchor, dismissed an "
     "interstitial itself, and read a value out of an iframe."),
    ("The discovered artifact replays", "It passes the same 8 scenarios as the hand-written "
     "one — and the approval gate refused it three times first, each refusal a real defect."),
    ("Two architectural controls", "No module in cua/ imports a model SDK. Nothing but "
     "surface.py imports Playwright. Both are enforced by tests."),
]
gaps = [
    ("Masking is off during discovery — on purpose",
     "The demo app holds synthetic members, so masking there protects nothing and can only "
     "cost accuracy. Built and tested; switching it on is a flag, and belongs with the first "
     "environment that may hold real data. Replay evidence has been masked throughout."),
    ("Screenshots are a deny-list",
     "The one channel with a partial guarantee, and recorded as the residual risk."),
    ("One held-out condition",
     "MAX_ACCOUNTS_REACHED (member 33333) is deliberately unhandled: it stays failed-with-"
     "unknown-state, and the verification check confirms the commit did not take effect."),
    ("One order-dependent test",
     "The operator-handoff test passes alone and fails in the full suite — the session-expiry "
     "scenario fires once per member, and an earlier test consumes it. A fixture issue, not an "
     "engine one, but it is not fixed yet."),
]

text(s, 0.62, 1.60, 5.8, 0.3, "PROVEN ON A REAL RUN", size=10.5, color=GREEN, bold=True)
for i, (title, body) in enumerate(proven):
    y = 1.98 + i * 1.22
    box(s, 0.62, y, 5.85, 1.08, "", fill=GREEN_BG, line=None)
    text(s, 0.80, y + 0.09, 5.5, 0.25, title, size=11, color=GREEN, bold=True)
    text(s, 0.80, y + 0.34, 5.5, 0.7, body, size=9.5, color=INK, spacing=2)

text(s, 6.87, 1.60, 5.8, 0.3, "KNOWN GAPS, NAMED", size=10.5, color=RED, bold=True)
for i, (title, body) in enumerate(gaps):
    y = 1.98 + i * 1.22
    box(s, 6.87, y, 5.85, 1.08, "", fill=RED_BG, line=None)
    text(s, 7.05, y + 0.09, 5.5, 0.25, title, size=11, color=RED, bold=True)
    text(s, 7.05, y + 0.34, 5.5, 0.7, body, size=9.5, color=INK, spacing=2)

footnote(s, "NOTES.md carries the full build log: what broke, why, and what changed as a result.")

prs.save(OUT)
print(f"wrote {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
